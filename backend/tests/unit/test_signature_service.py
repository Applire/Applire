# Copyright (C) 2026 Tobias Rosenbaum
# SPDX-License-Identifier: AGPL-3.0-or-later

"""#359 — signature service: storage, render-time resolution, format_place_date,
the two user_settings toggles, the DOCX writers' signature block, the
retention worker's orphan-scan reference, and migration 0065.

Covers ``applire.services.signature`` end to end plus the seams it touches:
``applire.routers.settings.get_settings``/``update_settings`` (the two
booleans), ``applire.services.office_export.cv_docx``/``letter_docx`` (the
signature block), ``applire.retention.worker._scan_orphan_files`` (the
``user_settings.signature_path`` reference), and ``alembic/versions/
0065_signature.py`` (read, not executed).

The HTTP surface (``routers/signature.py``'s POST/GET/DELETE) is tested
separately in ``tests/unit/test_signature_endpoints.py`` — this file stays at
the service layer, calling ``applire.services.signature`` functions and the
other modules above directly.

House style follows ``tests/unit/test_photo_service.py`` (in-memory SQLite via
``async_sessionmaker(expire_on_commit=False)``, ``LocalStorageProvider`` over
a temp dir) and ``tests/unit/test_retention_worker.py`` (the orphan-scan
fixtures, read there and re-created here against the real ORM models rather
than that file's raw-DDL SQLite harness — the ``UserSettings`` model already
carries ``signature_path`` via migration 0065, so no parallel DDL is needed).
"""
from __future__ import annotations

import io
import re
import time
import uuid
from datetime import date
from pathlib import Path

import pytest
import pytest_asyncio
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

# ---------------------------------------------------------------------------
# Shared fixtures
# ---------------------------------------------------------------------------


@pytest_asyncio.fixture
async def db():
    """In-memory SQLite session with every applire table (FK resolution)."""
    from applire.db.session import Base
    import applire.models.user  # noqa: F401
    import applire.models.profile  # noqa: F401
    import applire.models.job  # noqa: F401
    import applire.models.gap  # noqa: F401
    import applire.models.cv  # noqa: F401
    import applire.models.session  # noqa: F401
    import applire.models.flow  # noqa: F401
    import applire.models.uploads  # noqa: F401
    import applire.models.application  # noqa: F401
    import applire.models.color_profile  # noqa: F401
    import applire.models.company  # noqa: F401
    import applire.models.user_settings  # noqa: F401
    import applire.models.cover_letter  # noqa: F401

    engine = create_async_engine("sqlite+aiosqlite:///:memory:", echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    Session = async_sessionmaker(engine, expire_on_commit=False)
    async with Session() as session:
        yield session
    await engine.dispose()


@pytest.fixture
def storage(tmp_path):
    from applire.storage.local import LocalStorageProvider

    return LocalStorageProvider(str(tmp_path))


async def _make_user(db) -> uuid.UUID:
    """A real `users` row — `upload_signature`/`delete_signature`/
    `get_signature_bytes` all validate the user exists via `_get_user` before
    touching `user_settings` (which is keyed to the CE stub user regardless
    of which user_id was passed — ADR-022, single user in CE)."""
    from applire.models.user import User

    user_id = uuid.uuid4()
    db.add(User(id=user_id, email=f"{user_id}@example.de"))
    await db.commit()
    return user_id


async def _settings_row(db):
    from applire.models.user_settings import UserSettings
    from applire.services.color_detection import _CE_STUB_USER_ID

    result = await db.execute(
        select(UserSettings).where(UserSettings.user_id == _CE_STUB_USER_ID)
    )
    return result.scalar_one_or_none()


_PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"fake-signature-payload-not-a-real-png"


class _SaveFailsStorage:
    """Wraps a real StorageProvider but makes `save` fail — for pinning the
    save-before-delete ordering (services/signature.py's `upload_signature`,
    mirroring services/photo.py's own reasoning)."""

    def __init__(self, inner):
        self._inner = inner

    async def save(self, file_bytes, filename):
        raise OSError("disk full (simulated)")

    async def delete(self, file_path):
        return await self._inner.delete(file_path)

    async def read(self, file_path):
        return await self._inner.read(file_path)


class _RaisesFileNotFoundOnDelete:
    """A StorageProvider whose `delete` always raises FileNotFoundError —
    LocalStorageProvider.delete() swallows that internally, so this is the
    only way to actually exercise `delete_signature`'s own
    `except FileNotFoundError: pass` branch (signature.py:180-185)."""

    async def save(self, file_bytes, filename):
        raise NotImplementedError

    async def delete(self, file_path):
        raise FileNotFoundError(file_path)

    async def read(self, file_path):
        raise NotImplementedError


class _UnreachableStorage:
    """A StorageProvider that fails the test if any method is ever called —
    proves a codepath genuinely never touches storage, not just that its
    outcome happens to match (`Controls that don't fire` discipline)."""

    async def save(self, file_bytes, filename):
        raise AssertionError("storage.save() must not be called")

    async def delete(self, file_path):
        raise AssertionError("storage.delete() must not be called")

    async def read(self, file_path):
        raise AssertionError("storage.read() must not be called")


# ---------------------------------------------------------------------------
# upload_signature / delete_signature / get_signature_bytes — storage layer
# (properties 1-6)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_upload_signature_stores_bytes_sets_path_and_creates_settings_row(db, storage):
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    assert await _settings_row(db) is None, "no settings row must exist before the upload"

    result = await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
        db=db, storage=storage,
    )

    assert set(result.keys()) == {"signature_url"}
    row = await _settings_row(db)
    assert row is not None, "a signature upload must create the settings row itself"
    assert row.signature_path == result["signature_url"]
    assert Path(row.signature_path).read_bytes() == _PNG_BYTES


@pytest.mark.asyncio
async def test_upload_signature_does_not_require_a_profile(db, storage):
    """Unlike the photo (which refuses with 'import a CV first'), a signature
    upload must not require any MasterProfile row to exist."""
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    # No MasterProfile row is ever created in this test.
    result = await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
        db=db, storage=storage,
    )
    assert result["signature_url"] is not None


@pytest.mark.asyncio
async def test_reupload_replaces_old_file_and_deletes_it(db, storage):
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    first = await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
        db=db, storage=storage,
    )
    old_path = first["signature_url"]
    assert Path(old_path).exists()

    second = await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES + b"-v2", content_type="image/jpeg",
        db=db, storage=storage,
    )
    new_path = second["signature_url"]

    assert new_path != old_path
    assert Path(new_path).exists()
    assert not Path(old_path).exists(), "the old file must be deleted on replace"

    row = await _settings_row(db)
    assert row.signature_path == new_path


@pytest.mark.asyncio
async def test_upload_failure_on_save_leaves_old_path_intact(db, storage):
    """Save-before-delete: if the new save fails, the user keeps the
    signature they already had."""
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    first = await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
        db=db, storage=storage,
    )
    old_path = first["signature_url"]

    failing_storage = _SaveFailsStorage(storage)
    with pytest.raises(OSError):
        await upload_signature(
            user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
            db=db, storage=failing_storage,
        )

    row = await _settings_row(db)
    assert row.signature_path == old_path
    assert Path(old_path).exists()


@pytest.mark.asyncio
async def test_upload_rejects_unsupported_content_type(db, storage):
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    with pytest.raises(ValueError):
        await upload_signature(
            user_id=user_id, file_bytes=_PNG_BYTES, content_type="application/pdf",
            db=db, storage=storage,
        )
    assert await _settings_row(db) is None, "a rejected upload must write nothing"


@pytest.mark.asyncio
async def test_upload_rejects_empty_file(db, storage):
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    with pytest.raises(ValueError):
        await upload_signature(
            user_id=user_id, file_bytes=b"", content_type="image/png",
            db=db, storage=storage,
        )
    assert await _settings_row(db) is None


@pytest.mark.asyncio
async def test_upload_rejects_oversized_file(db, storage):
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    oversized = b"x" * (2 * 1024 * 1024 + 1)
    with pytest.raises(ValueError):
        await upload_signature(
            user_id=user_id, file_bytes=oversized, content_type="image/png",
            db=db, storage=storage,
        )
    assert await _settings_row(db) is None


@pytest.mark.asyncio
async def test_delete_signature_removes_file_and_clears_path(db, storage):
    from applire.services.signature import delete_signature, upload_signature

    user_id = await _make_user(db)
    result = await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
        db=db, storage=storage,
    )
    path = result["signature_url"]

    await delete_signature(user_id=user_id, db=db, storage=storage)

    assert not Path(path).exists()
    row = await _settings_row(db)
    assert row.signature_path is None


@pytest.mark.asyncio
async def test_delete_signature_is_noop_when_nothing_on_file(db, storage):
    from applire.services.signature import delete_signature

    user_id = await _make_user(db)
    # No upload at all — no settings row exists yet.
    await delete_signature(user_id=user_id, db=db, storage=storage)  # must not raise
    assert await _settings_row(db) is None, "a no-op delete must not create a row"


@pytest.mark.asyncio
async def test_delete_signature_clears_path_even_if_file_already_gone(db, storage):
    """The FileNotFoundError branch: delete_signature's own try/except, not
    LocalStorageProvider's (which already swallows FileNotFoundError itself
    and so cannot exercise this branch)."""
    from applire.models.user_settings import UserSettings
    from applire.services.color_detection import _CE_STUB_USER_ID
    from applire.services.signature import delete_signature

    user_id = await _make_user(db)
    db.add(UserSettings(user_id=_CE_STUB_USER_ID, signature_path="/does/not/matter.png"))
    await db.commit()

    await delete_signature(
        user_id=user_id, db=db, storage=_RaisesFileNotFoundOnDelete()
    )  # must not raise despite storage.delete() raising

    row = await _settings_row(db)
    assert row.signature_path is None


@pytest.mark.asyncio
async def test_get_signature_bytes_returns_bytes_and_mime(db, storage):
    from applire.services.signature import get_signature_bytes, upload_signature

    user_id = await _make_user(db)
    await upload_signature(
        user_id=user_id, file_bytes=_PNG_BYTES, content_type="image/png",
        db=db, storage=storage,
    )

    raw, mime = await get_signature_bytes(user_id=user_id, db=db, storage=storage)
    assert raw == _PNG_BYTES
    assert mime == "image/png"


@pytest.mark.asyncio
async def test_get_signature_bytes_raises_lookuperror_when_none_on_file(db, storage):
    from applire.services.signature import get_signature_bytes

    user_id = await _make_user(db)
    with pytest.raises(LookupError):
        await get_signature_bytes(user_id=user_id, db=db, storage=storage)


# ---------------------------------------------------------------------------
# resolve_signature_data_uri / resolve_signature_bytes — render-time
# resolution, the toggle (properties 7-12)
# ---------------------------------------------------------------------------


async def _upload(db, storage, content_type="image/png", file_bytes=_PNG_BYTES):
    from applire.services.signature import upload_signature

    user_id = await _make_user(db)
    return await upload_signature(
        user_id=user_id, file_bytes=file_bytes, content_type=content_type,
        db=db, storage=storage,
    )


@pytest.mark.asyncio
async def test_resolve_letter_data_uri_is_present_by_default(db, storage, monkeypatch):
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_data_uri

    await _upload(db, storage)
    result = await resolve_signature_data_uri(db, document="letter")
    assert result is not None
    assert result.startswith("data:image/png;base64,")


@pytest.mark.asyncio
async def test_resolve_letter_data_uri_is_none_when_toggle_off(db, storage, monkeypatch):
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_data_uri

    await _upload(db, storage)
    row = await _settings_row(db)
    row.signature_in_letter = False
    await db.commit()

    result = await resolve_signature_data_uri(db, document="letter")
    assert result is None


@pytest.mark.asyncio
async def test_resolve_cv_data_uri_is_none_by_default(db, storage, monkeypatch):
    """The founder's default is CV OFF (F-0)."""
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_data_uri

    await _upload(db, storage)
    row = await _settings_row(db)
    assert row.signature_in_cv is False, "sanity: the stored default really is False"

    result = await resolve_signature_data_uri(db, document="cv")
    assert result is None


@pytest.mark.asyncio
async def test_resolve_cv_data_uri_present_once_toggle_enabled(db, storage, monkeypatch):
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_data_uri

    await _upload(db, storage)
    row = await _settings_row(db)
    row.signature_in_cv = True
    await db.commit()

    result = await resolve_signature_data_uri(db, document="cv")
    assert result is not None
    assert result.startswith("data:image/png;base64,")


@pytest.mark.asyncio
async def test_resolve_data_uri_none_for_both_kinds_with_no_settings_row(db):
    """No row at all == the toggle question never arises for either document
    kind, and storage must never even be consulted."""
    from applire.services.signature import resolve_signature_data_uri

    assert await _settings_row(db) is None

    for kind in ("letter", "cv"):
        import applire.storage as storage_module

        original = storage_module.get_storage
        storage_module.get_storage = lambda: _UnreachableStorage()
        try:
            result = await resolve_signature_data_uri(db, document=kind)
        finally:
            storage_module.get_storage = original
        assert result is None, f"document={kind!r}"


@pytest.mark.asyncio
async def test_resolve_data_uri_missing_file_returns_none_not_raise(db, storage, monkeypatch):
    """A deleted asset must not turn a download into a 500."""
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_data_uri

    await _upload(db, storage)
    row = await _settings_row(db)
    Path(row.signature_path).unlink()  # file gone, DB still points at it

    result = await resolve_signature_data_uri(db, document="letter")  # must not raise
    assert result is None


@pytest.mark.asyncio
async def test_resolve_signature_bytes_roundtrips(db, storage, monkeypatch):
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_bytes

    await _upload(db, storage)
    result = await resolve_signature_bytes(db, document="letter")
    assert result == _PNG_BYTES


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "content_type,expected_mime",
    [("image/png", "image/png"), ("image/jpeg", "image/jpeg"), ("image/webp", "image/webp")],
)
async def test_resolve_data_uri_mime_follows_stored_suffix(
    db, storage, monkeypatch, content_type, expected_mime
):
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)
    from applire.services.signature import resolve_signature_data_uri

    await _upload(db, storage, content_type=content_type)
    result = await resolve_signature_data_uri(db, document="letter")
    assert result.startswith(f"data:{expected_mime};base64,")


# ---------------------------------------------------------------------------
# format_place_date (property 13)
# ---------------------------------------------------------------------------


def test_format_place_date_german():
    from applire.services.signature import format_place_date

    assert format_place_date("Berlin", "de", date(2026, 9, 11)) == "Berlin, 11. September 2026"


def test_format_place_date_english():
    from applire.services.signature import format_place_date

    assert format_place_date("Berlin", "en", date(2026, 9, 11)) == "Berlin, 11 September 2026"


@pytest.mark.parametrize("location", [None, "", "   "])
def test_format_place_date_drops_the_whole_ort_part_when_location_is_absent(location):
    """A fabricated city on a signed document is a truthfulness defect, not a
    cosmetic one — assert no placeholder and no dangling comma."""
    from applire.services.signature import format_place_date

    result = format_place_date(location, "de", date(2026, 9, 11))
    assert result == "11. September 2026"
    assert "," not in result
    assert not result.startswith(",")


# ---------------------------------------------------------------------------
# Settings contract — get_settings / update_settings / the two schemas
# (properties 14-16)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_get_settings_defaults_letter_on_cv_off_on_a_fresh_db(db):
    from applire.routers.settings import get_settings

    result = await get_settings(db)
    assert result["signature_in_letter"] is True
    assert result["signature_in_cv"] is False


@pytest.mark.asyncio
async def test_update_settings_writes_each_toggle_independently(db):
    from applire.routers.settings import update_settings

    r1 = await update_settings(db, signature_in_letter=False)
    assert r1["signature_in_letter"] is False
    assert r1["signature_in_cv"] is False, "untouched, still its own default"

    r2 = await update_settings(db, signature_in_cv=True)
    assert r2["signature_in_cv"] is True
    assert r2["signature_in_letter"] is False, "untouched by the second PATCH"


@pytest.mark.asyncio
async def test_update_settings_omitting_both_toggles_leaves_them_untouched(db):
    from applire.routers.settings import update_settings

    await update_settings(db, signature_in_letter=False, signature_in_cv=True)
    result = await update_settings(db, ui_language="de")  # both toggles omitted (None)

    assert result["signature_in_letter"] is False
    assert result["signature_in_cv"] is True


def test_settings_response_and_patch_request_carry_the_signature_fields():
    from applire.routers.settings import SettingsPatchRequest, SettingsResponse

    assert "signature_in_letter" in SettingsResponse.model_fields
    assert "signature_in_cv" in SettingsResponse.model_fields
    assert SettingsResponse.model_fields["signature_in_letter"].default is True
    assert SettingsResponse.model_fields["signature_in_cv"].default is False

    assert "signature_in_letter" in SettingsPatchRequest.model_fields
    assert "signature_in_cv" in SettingsPatchRequest.model_fields
    assert SettingsPatchRequest.model_fields["signature_in_letter"].default is None
    assert SettingsPatchRequest.model_fields["signature_in_cv"].default is None


# ---------------------------------------------------------------------------
# DOCX writers — the signature block (properties 17-19)
# ---------------------------------------------------------------------------


def _real_png_bytes(color=(30, 60, 120)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (40, 20), color).save(buf, format="PNG")
    return buf.getvalue()


def _minimal_tailored_cv():
    from applire.schemas.cv import TailoredCVData, TailoredContact

    return TailoredCVData(
        contact=TailoredContact(name="Anna Bauer"), summary="MARKER_SUMMARY_TEXT"
    )


def _minimal_letter_data():
    from applire.schemas.cover_letter import LetterBody, LetterData, LetterSignature

    return LetterData(
        body=LetterBody(paragraphs=["MARKER_BODY_TEXT"]),
        signature=LetterSignature(closing="MARKER_CLOSING_TEXT", name="MARKER_SIGNATURE_NAME"),
    )


def _open_docx(docx_bytes: bytes):
    from docx import Document as open_docx

    return open_docx(io.BytesIO(docx_bytes))


def test_render_cv_docx_signature_block_present_with_bytes():
    from applire.services.office_export.cv_docx import render_cv_docx

    result = render_cv_docx(
        _minimal_tailored_cv(), lang="de", accent_color="#1a3a5c",
        signature_bytes=_real_png_bytes(),
        signature_place_date="Berlin, 11. September 2026",
    )
    document = _open_docx(result)
    text = "\n".join(p.text for p in document.paragraphs)

    assert "Berlin, 11. September 2026" in text
    assert len(document.inline_shapes) == 1


def test_render_cv_docx_signature_block_absent_without_bytes():
    from applire.services.office_export.cv_docx import render_cv_docx

    result = render_cv_docx(
        _minimal_tailored_cv(), lang="de", accent_color="#1a3a5c",
        signature_bytes=None, signature_place_date="Berlin, 11. September 2026",
    )
    document = _open_docx(result)
    text = "\n".join(p.text for p in document.paragraphs)

    assert "Berlin, 11. September 2026" not in text
    assert len(document.inline_shapes) == 0


def test_render_letter_docx_embeds_signature_between_closing_and_name():
    from applire.services.office_export.letter_docx import render_letter_docx

    letter = _minimal_letter_data()
    with_sig = _open_docx(
        render_letter_docx(letter, lang="de", accent_color="#1a3a5c", signature_bytes=_real_png_bytes())
    )
    without_sig = _open_docx(
        render_letter_docx(letter, lang="de", accent_color="#1a3a5c", signature_bytes=None)
    )

    assert len(with_sig.inline_shapes) == 1
    assert len(without_sig.inline_shapes) == 0

    def _gap(document):
        paras = document.paragraphs
        close_idx = next(i for i, p in enumerate(paras) if "MARKER_CLOSING_TEXT" in p.text)
        name_idx = next(i for i, p in enumerate(paras) if p.text == "MARKER_SIGNATURE_NAME")
        return name_idx - close_idx

    assert _gap(without_sig) == 1, "closing then name, adjacent, when there is no signature"
    assert _gap(with_sig) == 2, "closing, then the picture's own paragraph, then name"


def test_render_cv_docx_unreadable_signature_bytes_do_not_raise():
    """A webp photo once produced an HTTP 500 on every .docx download
    (_render_contact's identical guard) — this is the signature's version of
    that regression."""
    from applire.services.office_export.cv_docx import render_cv_docx

    result = render_cv_docx(
        _minimal_tailored_cv(), lang="de", accent_color="#1a3a5c",
        signature_bytes=b"not an image at all",
        signature_place_date="Berlin, 11. September 2026",
    )
    assert result[:2] == b"PK"
    text = "\n".join(p.text for p in _open_docx(result).paragraphs)
    assert "MARKER_SUMMARY_TEXT" in text, "the rest of the document must still render"


def test_render_letter_docx_unreadable_signature_bytes_do_not_raise():
    from applire.services.office_export.letter_docx import render_letter_docx

    letter = _minimal_letter_data()
    result = render_letter_docx(
        letter, lang="de", accent_color="#1a3a5c", signature_bytes=b"not an image at all"
    )
    assert result[:2] == b"PK"
    text = "\n".join(p.text for p in _open_docx(result).paragraphs)
    assert "MARKER_BODY_TEXT" in text
    assert "MARKER_SIGNATURE_NAME" in text


# ---------------------------------------------------------------------------
# Retention — the orphan scan reads user_settings.signature_path (#152 /
# SF-PROFILE.5, ADR-088) (property 20)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_orphan_scan_keeps_referenced_signature_and_deletes_unreferenced_file(
    db, tmp_path, monkeypatch
):
    """tests/unit/test_retention_worker.py already covers the general orphan
    scan with its own raw-DDL SQLite harness; this uses the real ORM models
    (the `db` fixture here already creates `uploads`/`master_profiles`/
    `user_settings` via Base.metadata.create_all, so _scan_orphan_files's
    three SELECTs all succeed against real, if empty, tables)."""
    import os

    from applire.models.user_settings import UserSettings
    from applire.retention.worker import _scan_orphan_files
    from applire.services.color_detection import _CE_STUB_USER_ID
    from applire.storage.local import LocalStorageProvider

    storage = LocalStorageProvider(str(tmp_path))
    monkeypatch.setattr("applire.storage.get_storage", lambda: storage)

    sig_path = await storage.save(b"referenced-signature-bytes", "signature.png")
    orphan_path = await storage.save(b"unreferenced-bytes", "orphan.png")

    old = time.time() - 3 * 3600  # past the (default 1h) grace period
    os.utime(sig_path, (old, old))
    os.utime(orphan_path, (old, old))

    db.add(UserSettings(user_id=_CE_STUB_USER_ID, signature_path=sig_path))
    await db.commit()

    deleted = await _scan_orphan_files(db)

    assert deleted == 1
    assert Path(sig_path).exists(), (
        "a signature referenced by user_settings.signature_path must survive "
        "the orphan scan"
    )
    assert not Path(orphan_path).exists(), (
        "an unreferenced file next to it must still be reclaimed"
    )


# ---------------------------------------------------------------------------
# Migration 0065 — read the module, never run alembic (property 21)
# ---------------------------------------------------------------------------


def _load_migration_0065():
    import importlib.util

    path = Path(__file__).resolve().parents[2] / "alembic" / "versions" / "0065_signature.py"
    spec = importlib.util.spec_from_file_location(
        "_migration_0065_signature_under_test", path
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_migration_0065_revision_chain_and_touches_exactly_three_columns():
    import inspect

    module = _load_migration_0065()
    assert module.revision == "0065"
    assert module.down_revision == "0064"

    upgrade_src = inspect.getsource(module.upgrade)
    downgrade_src = inspect.getsource(module.downgrade)

    added = re.findall(
        r'op\.add_column\(\s*"user_settings",\s*sa\.Column\(\s*"(\w+)"', upgrade_src
    )
    dropped = re.findall(
        r'op\.drop_column\(\s*"user_settings",\s*"(\w+)"\)', downgrade_src
    )

    expected = {"signature_path", "signature_in_letter", "signature_in_cv"}
    assert set(added) == expected and len(added) == 3
    assert set(dropped) == expected and len(dropped) == 3
